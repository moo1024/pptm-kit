"""그라데이션 테두리 · 그라데이션 글씨 · 빛 번짐 — 렌더가 끝난 deck.pptx 에 덧입히는 후처리.

XML 은 PowerPoint 실물 파일에서 뗐다.
  · 테두리  slide10  <a:ln w="13970"><a:gradFill> FF9D2A·FD7ABB·7F66FF·2F6DEA <a:lin ang="5400000"/>
  · 원 고리  slide3   <a:ln w="25400"><a:gradFill> … <a:path path="circle">
  · 글씨    slide1   <a:rPr><a:gradFill> … <a:lin ang="0" scaled="1"/> (+ 짙은 outerShdw)
  · 빛 번짐  slide3   <a:glow rad="63500"> bg1 alpha 15%

렌더러(gyoan.py)를 고치지 않으려고 **이름표로 도형을 찾는다**. 스펙 아이템에
`"fx": {...}` 를 주고 `name`(또는 `morph_id` — 도형 이름이 `!!id` 가 된다)을 같이 단다.

    {"type": "card", ..., "name": "tobe", "fx": {"line": {"stops": ["FF9D2A", "7F66FF"], "w": 1.1, "ang": 90}}}
    {"type": "text", ..., "name": "kw",   "fx": {"text": {"stops": ["FF9D2A", "FD7ABB", "7F66FF"]}}}
    {"type": "card", ..., "name": "ring", "fx": {"glow": {"color": "7F66FF", "rad": 14, "alpha": 0.35}}}

`fx.text` 는 `runs` 의 몇 번째 런만 고를 수 있다 — `{"stops": …, "runs": [1]}` (줄 구분 없이 런 순번).

**LibreOffice 는 선 그라데이션을 단색으로, 글씨 그라데이션을 첫 색으로 그린다.**
그래서 이 모듈을 쓴 덱의 눈검수는 PowerPoint 로 뜬 그림(`export_ppt_png`)으로 한다.

실행:  python3 -m pptm.design.gradfx <slug>      (pptm rendernative 다음, pptm export 앞)
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]


def _grad_el(stops: list[str], ang: float = 0.0, path: str | None = None,
             alphas: list[float] | None = None) -> etree._Element:
    """<a:gradFill> — stops 를 고르게 펼친다. path="circle" 이면 원형(고리)."""
    g = etree.Element(qn("a:gradFill"), {"flip": "none", "rotWithShape": "1"})
    lst = etree.SubElement(g, qn("a:gsLst"))
    n = len(stops)
    for i, hx in enumerate(stops):
        gs = etree.SubElement(lst, qn("a:gs"), {"pos": str(int(round(100000 * i / max(1, n - 1))))})
        c = etree.SubElement(gs, qn("a:srgbClr"), {"val": hx.lstrip("#").upper()})
        if alphas is not None and alphas[i] < 1:
            etree.SubElement(c, qn("a:alpha"), {"val": str(int(alphas[i] * 100000))})
    if path:
        p = etree.SubElement(g, qn("a:path"), {"path": path})
        etree.SubElement(p, qn("a:fillToRect"), {"r": "100000", "b": "100000"})
        etree.SubElement(g, qn("a:tileRect"), {"l": "-100000", "t": "-100000"})
    else:
        etree.SubElement(g, qn("a:lin"), {"ang": str(int(ang * 60000)), "scaled": "1"})
        etree.SubElement(g, qn("a:tileRect"))
    return g


_FILL_TAGS = ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill")


def line_grad(sp, stops, w=1.1, ang=90.0, path=None, alphas=None, dash=None):
    """도형 외곽선을 그라데이션으로 — spPr 의 a:ln 을 통째로 바꾼다."""
    spPr = sp.find(qn("p:spPr"))
    old = spPr.find(qn("a:ln"))
    ln = etree.Element(qn("a:ln"), {"w": str(int(w * 12700))})
    ln.append(_grad_el(stops, ang, path, alphas))
    if dash:
        etree.SubElement(ln, qn("a:prstDash"), {"val": dash})
    if old is not None:
        spPr.replace(old, ln)
    else:   # 스키마 순서: xfrm·geom·fill 다음, effectLst 앞
        eff = spPr.find(qn("a:effectLst"))
        if eff is not None:
            eff.addprevious(ln)
        else:
            spPr.append(ln)
    # p:style 의 lnRef 가 남아 있으면 PowerPoint 가 테마 선을 덧칠할 수 있다 — 렌더러가 이미 지우지만 한 번 더
    st = sp.find(qn("p:style"))
    if st is not None:
        sp.remove(st)


def text_grad(sp, stops, ang=0.0, runs=None, shadow=None):
    """글자 채움을 그라데이션으로 — rPr 의 solidFill 을 바꾼다. runs=None 이면 전부."""
    rprs = []
    for r in sp.iter(qn("a:r")):
        rp = r.find(qn("a:rPr"))
        if rp is None:
            rp = etree.SubElement(r, qn("a:rPr"))
            r.insert(0, rp)
        rprs.append(rp)
    pick = range(len(rprs)) if runs is None else runs
    for i in pick:
        if i >= len(rprs):
            continue
        rp = rprs[i]
        for tag in _FILL_TAGS:
            for el in rp.findall(qn(tag)):
                rp.remove(el)
        g = _grad_el(stops, ang)
        # 스키마 순서: ln? → fill → effectLst … → latin/ea/cs
        ln = rp.find(qn("a:ln"))
        if ln is not None:
            ln.addnext(g)
        else:
            rp.insert(0, g)
        if shadow:   # 그라데이션 글씨 밑의 짙은 번짐 그림자
            eff = rp.find(qn("a:effectLst"))
            if eff is None:
                eff = etree.Element(qn("a:effectLst"))
                g.addnext(eff)
            o = etree.SubElement(eff, qn("a:outerShdw"), {
                "blurRad": str(int(shadow.get("blur", 14) * 12700)), "algn": "ctr", "rotWithShape": "0"})
            c = etree.SubElement(o, qn("a:srgbClr"), {"val": shadow.get("color", "0C0E1A")})
            etree.SubElement(c, qn("a:alpha"), {"val": str(int(shadow.get("alpha", 0.84) * 100000))})


def fill_grad(sp, stops, ang=0.0, path=None, alphas=None):
    """면 채움을 여러 색 그라데이션으로 — 렌더러 grad 는 두 색·세로뿐이라 알약 배지용."""
    spPr = sp.find(qn("p:spPr"))
    for tag in _FILL_TAGS:
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)
    g = _grad_el(stops, ang, path, alphas)
    geom = spPr.find(qn("a:prstGeom"))
    if geom is None:
        geom = spPr.find(qn("a:custGeom"))
    geom.addnext(g)


def glow(sp, color="FFFFFF", rad=5.0, alpha=0.15):
    """도형 빛 번짐 — spPr/effectLst/glow."""
    spPr = sp.find(qn("p:spPr"))
    eff = spPr.find(qn("a:effectLst"))
    if eff is None:
        eff = etree.SubElement(spPr, qn("a:effectLst"))
    for el in eff.findall(qn("a:glow")):
        eff.remove(el)
    g = etree.Element(qn("a:glow"), {"rad": str(int(rad * 12700))})
    c = etree.SubElement(g, qn("a:srgbClr"), {"val": color.lstrip("#").upper()})
    etree.SubElement(c, qn("a:alpha"), {"val": str(int(alpha * 100000))})
    eff.insert(0, g)   # 스키마 순서: blur → fillOverlay → glow → innerShdw → outerShdw …


# ── 차트 — forest(점 + 신뢰구간) ──────────────────────────────────────────
# 렌더러 chart 는 line_markers 까지만 있다. 선을 숨기고 점마다 색을 주고 사용자 지정 오차 막대
# (c:errBars errValType=cust)를 붙이면 편집 가능한 네이티브 forest 가 된다. 값은 엑셀 백업에 그대로 남는다.
def _c(tag):
    return qn("c:" + tag)


def _solid(parent, hx, alpha=None):
    f = etree.SubElement(parent, qn("a:solidFill"))
    c = etree.SubElement(f, qn("a:srgbClr"), {"val": hx.lstrip("#").upper()})
    if alpha is not None:
        etree.SubElement(c, qn("a:alpha"), {"val": str(int(alpha * 100000))})
    return f


def _ser_insert(ser, el, before=("dLbls", "trendline", "errBars", "cat", "val", "smooth", "extLst")):
    for tag in before:
        ref = ser.find(_c(tag))
        if ref is not None:
            ref.addprevious(el)
            return
    ser.append(el)


def _numlit(vals):
    nl = etree.Element(_c("numLit"))
    etree.SubElement(nl, _c("formatCode")).text = "General"
    etree.SubElement(nl, _c("ptCount"), {"val": str(len(vals))})
    for i, v in enumerate(vals):
        if v is None:
            continue
        pt = etree.SubElement(nl, _c("pt"), {"idx": str(i)})
        etree.SubElement(pt, _c("v")).text = repr(float(v))
    return nl


def chart_fx(gf, spec: dict) -> None:
    """spec 키 — hide_line:[계열] · points:{계열:{colors,size,symbol}} ·
    errbars:{계열:{plus,minus,color,w}} · line:{계열:{color,w,dash}} · no_marker:[계열] · major_unit ·
    keep_negative(막대 음수를 흰색으로 뒤집지 않음) · tick_low(항목 라벨을 축 바닥에 — 음수 막대와 안 겹침)"""
    cs = gf.chart._chartSpace
    if spec.get("keep_negative"):
        # 점별 색(dPt)에 invertIfNegative 가 없으면 PowerPoint 는 음수 막대를 흰색으로 뒤집는다(2026-10-09 실측)
        for d in cs.iter(_c("dPt")):
            if d.find(_c("invertIfNegative")) is None:
                d.find(_c("idx")).addnext(etree.Element(_c("invertIfNegative"), {"val": "0"}))
        for ser in cs.iter(_c("ser")):
            for el in ser.findall(_c("invertIfNegative")):
                el.set("val", "0")
    if spec.get("tick_low"):
        for ax in cs.iter(_c("catAx")):
            t = ax.find(_c("tickLblPos"))
            if t is None:
                t = etree.Element(_c("tickLblPos"))
                ref = next(e for e in (ax.find(_c(k)) for k in ("spPr", "txPr", "crossAx")) if e is not None)
                ref.addprevious(t)
            t.set("val", "low")
    if spec.get("major_unit"):
        for ax in cs.iter(_c("valAx")):
            for el in ax.findall(_c("majorUnit")):
                ax.remove(el)
            mu = etree.Element(_c("majorUnit"), {"val": str(spec["major_unit"])})
            ref = None
            for tag in ("minorUnit", "dispUnits", "extLst"):
                ref = ax.find(_c(tag))
                if ref is not None:
                    break
            if ref is not None:
                ref.addprevious(mu)
            else:
                ax.append(mu)
    sers = list(cs.iter(_c("ser")))
    for k in spec.get("hide_line", []):
        ser = sers[int(k)]
        sp = ser.find(_c("spPr"))
        if sp is None:
            sp = etree.Element(_c("spPr"))
            _ser_insert(ser, sp, ("marker", "dPt", "dLbls", "cat", "val", "smooth"))
        for el in sp.findall(qn("a:ln")):
            sp.remove(el)
        ln = etree.SubElement(sp, qn("a:ln"))
        etree.SubElement(ln, qn("a:noFill"))
    for k, ls in spec.get("line", {}).items():
        ser = sers[int(k)]
        sp = ser.find(_c("spPr"))
        for el in sp.findall(qn("a:ln")):
            sp.remove(el)
        ln = etree.SubElement(sp, qn("a:ln"), {"w": str(int(ls.get("w", 1.0) * 12700))})
        _solid(ln, ls["color"], ls.get("alpha"))
        if ls.get("dash"):
            etree.SubElement(ln, qn("a:prstDash"), {"val": ls["dash"]})
    for k in spec.get("no_marker", []):
        ser = sers[int(k)]
        m = ser.find(_c("marker"))
        if m is not None:
            ser.remove(m)
        m = etree.Element(_c("marker"))
        etree.SubElement(m, _c("symbol"), {"val": "none"})
        _ser_insert(ser, m, ("dPt", "dLbls", "cat", "val", "smooth"))
    for k, ps in spec.get("points", {}).items():
        ser = sers[int(k)]
        for el in ser.findall(_c("dPt")):
            ser.remove(el)
        for i, hx in enumerate(ps["colors"]):
            if not hx:
                continue
            d = etree.Element(_c("dPt"))
            etree.SubElement(d, _c("idx"), {"val": str(i)})
            m = etree.SubElement(d, _c("marker"))
            etree.SubElement(m, _c("symbol"), {"val": ps.get("symbol", "circle")})
            etree.SubElement(m, _c("size"), {"val": str(ps.get("size", 9))})
            msp = etree.SubElement(m, _c("spPr"))
            _solid(msp, hx)
            mln = etree.SubElement(msp, qn("a:ln"), {"w": str(int(ps.get("ring_w", 1.5) * 12700))})
            if ps.get("ring"):
                _solid(mln, ps["ring"])
            else:
                etree.SubElement(mln, qn("a:noFill"))
            etree.SubElement(d, _c("bubble3D"), {"val": "0"})
            _ser_insert(ser, d, ("dLbls", "trendline", "errBars", "cat", "val", "smooth"))
    for k, eb in spec.get("errbars", {}).items():
        ser = sers[int(k)]
        for el in ser.findall(_c("errBars")):
            ser.remove(el)
        e = etree.Element(_c("errBars"))
        etree.SubElement(e, _c("errDir"), {"val": "y"})
        etree.SubElement(e, _c("errBarType"), {"val": "both"})
        etree.SubElement(e, _c("errValType"), {"val": "cust"})
        etree.SubElement(e, _c("noEndCap"), {"val": "0" if eb.get("caps", True) else "1"})
        etree.SubElement(e, _c("plus")).append(_numlit(eb["plus"]))
        etree.SubElement(e, _c("minus")).append(_numlit(eb["minus"]))
        sp = etree.SubElement(e, _c("spPr"))
        ln = etree.SubElement(sp, qn("a:ln"), {"w": str(int(eb.get("w", 1.75) * 12700)), "cap": "rnd"})
        _solid(ln, eb.get("color", "9FA0A4"))
        _ser_insert(ser, e, ("cat", "val", "smooth", "extLst"))


def _apply_fx(sp, fx: dict) -> list[str]:
    done = []
    if fx.get("fill"):
        fill_grad(sp, **fx["fill"])
        done.append("fill")
    if fx.get("line"):
        line_grad(sp, **fx["line"])
        done.append("line")
    if fx.get("text"):
        text_grad(sp, **fx["text"])
        done.append("text")
    if fx.get("glow"):
        glow(sp, **fx["glow"])
        done.append("glow")
    return done


def _shape_names(item) -> list[str]:
    out = []
    if item.get("morph_id"):
        out.append(f"!!{item['morph_id']}")
    if item.get("name"):
        out.append(str(item["name"]))
    return out


def apply_deck(project_dir: Path, pptx: Path | None = None) -> list[str]:
    """스펙의 fx 를 deck.pptx 에 덧입힌다. 못 찾은 이름표는 오류로 멈춘다(조용히 통과 금지)."""
    from pptm import pages as pages_mod
    project_dir = Path(project_dir)
    pptx = pptx or project_dir / "output" / "deck.pptx"
    prs = Presentation(str(pptx))
    specs = [json.loads((d / "layout_spec.json").read_text())
             for d in pages_mod.page_dirs(project_dir) if (d / "layout_spec.json").exists()]
    log, missing = [], []
    for slide, spec in zip(prs.slides, specs):
        by_name = {}

        def walk(shapes):        # 묶음(grpSp) 안 도형도 찾는다 — group 키로 묶인 아이템
            for shp in shapes:
                by_name.setdefault(shp.name, shp)
                if shp.shape_type == 6:      # MSO_SHAPE_TYPE.GROUP
                    walk(shp.shapes)
        walk(slide.shapes)
        for item in spec.get("items", []):
            fx = item.get("fx")
            if not fx:
                continue
            shp = next((by_name[n] for n in _shape_names(item) if n in by_name), None)
            if shp is None:
                missing.append(f"p{spec.get('page')}: {_shape_names(item) or item.get('type')}")
                continue
            if fx.get("chart"):
                chart_fx(shp, fx["chart"])
                done = ["chart"]
            else:
                done = _apply_fx(shp._element, fx)
            log.append(f"p{spec.get('page')} {_shape_names(item)[0]}: {','.join(done)}")
    if missing:
        raise SystemExit("✗ fx 이름표를 못 찾음 — 아이템에 name 이나 morph_id 를 달아라:\n  "
                         + "\n  ".join(missing))
    prs.save(str(pptx))
    return log


if __name__ == "__main__":
    slug = sys.argv[1]
    for line in apply_deck(ROOT / "projects" / slug):
        print(line)
