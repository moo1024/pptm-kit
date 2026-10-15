"""3D 모델 · 영상 · 소리 아이템 — python-pptx 에 API 가 없거나 모자라 XML 을 직접 쓴다.

구조는 추측이 아니라 실물에서 뗀 그대로다(2026-10-08).
- 3D 모델  (`ppt/media/model3d1.glb`) — `mc:AlternateContent` 안
           Choice(Requires="am3d") `p:graphicFrame` > `am3d:model3d` + Fallback `p:pic`(정지 그림).
           관계 `http://schemas.microsoft.com/office/2017/06/relationships/model3d`, 형식 `model/gltf.binary`.
           회전 `am3d:trans/am3d:rot ax·ay·az`(1/60000도) · 카메라 `am3d:camera pos·up·lookAt·perspective fov`.
           장 사이 모핑이 회전·카메라·크기를 보간한다.
- 영상     `p:pic` + `a:videoFile r:link`(관계 video) +
           `p14:media r:embed`(관계 media) + `a:hlinkClick action="ppaction://media"` + 포스터 blip.
           재생은 timing 쪽(motion.py) — mainSeq `mediacall playFrom(0.0)` · 클릭 토글 interactiveSeq ·
           `p:video/p:cMediaNode`(음량·반복·멈추면 숨김).
- 소리     영상과 같은 골격에 `a:audioFile`(관계 audio). 5덱엔 없다 — OOXML 같은 자리라 같은 코드로 쓴다.
"""
from __future__ import annotations

import io
import json
import struct
from pathlib import Path

from lxml import etree
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI

NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NS_AM3D = "http://schemas.microsoft.com/office/drawing/2017/model3d"
NS_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
RT_MODEL3D = "http://schemas.microsoft.com/office/2017/06/relationships/model3d"
CT_GLB = "model/gltf.binary"

# 실물 기본값 — 카메라·조명·변환. 다른 glb 는 `trans` 로 덮는다.
CAMERA = {"pos": [0, 0, 81469193], "up": [0, 36000000, 0], "lookAt": [0, 0, 0], "fov": 2700000}
TRANS = {"meterPerModelUnit": [5000000, 1000000], "preTrans": [-3, 0, -3],
         "scale": [1, 1, 1], "postTrans": [0, 0, 0]}
LIGHTS = (
    '<am3d:ambientLight><am3d:clr><a:scrgbClr r="50000" g="50000" b="50000"/></am3d:clr>'
    '<am3d:illuminance n="500000" d="1000000"/></am3d:ambientLight>'
    '<am3d:ptLight rad="0"><am3d:clr><a:scrgbClr r="100000" g="75000" b="50000"/></am3d:clr>'
    '<am3d:intensity n="9765625" d="1000000"/><am3d:pos x="21959998" y="70920001" z="16344003"/></am3d:ptLight>'
    '<am3d:ptLight rad="0"><am3d:clr><a:scrgbClr r="40000" g="60000" b="95000"/></am3d:clr>'
    '<am3d:intensity n="12250000" d="1000000"/><am3d:pos x="-37964106" y="51130435" z="57631972"/></am3d:ptLight>'
    '<am3d:ptLight rad="0"><am3d:clr><a:scrgbClr r="86837" g="72700" b="100000"/></am3d:clr>'
    '<am3d:intensity n="3125000" d="1000000"/><am3d:pos x="-37739122" y="58056624" z="-34769649"/></am3d:ptLight>')
VIEWPORT_RATIO = 1.768   # 실측 viewportSz / 틀 폭(1.768~1.812)


def _next_shape_id(slide) -> int:
    ids = [int(x) for x in slide.shapes._spTree.xpath("//p:cNvPr/@id") if str(x).isdigit()]
    return max(ids + [1]) + 1


def _glb_part(slide, path: Path):
    """같은 glb 는 덱 안에서 파트 하나를 나눠 쓴다."""
    pkg = slide.part.package
    cache = pkg.__dict__.setdefault("_pm_glb", {})
    key = str(Path(path).resolve())
    if key not in cache:
        uri = pkg.next_partname("/ppt/media/model3d%d.glb")
        cache[key] = Part(PackURI(uri), CT_GLB, pkg, Path(path).read_bytes())
    return cache[key]


def glb_bbox(path: Path):
    """glb JSON 청크에서 POSITION 접근자 min/max — 일반 glb 의 중심·크기 추정용."""
    b = Path(path).read_bytes()
    n = struct.unpack("<I", b[12:16])[0]
    j = json.loads(b[20:20 + n])
    lo, hi = [1e30] * 3, [-1e30] * 3
    for m in j.get("meshes", []):
        for pr in m.get("primitives", []):
            a = j["accessors"][pr["attributes"]["POSITION"]]
            if "min" in a:
                lo = [min(x, y) for x, y in zip(lo, a["min"])]
                hi = [max(x, y) for x, y in zip(hi, a["max"])]
    return lo, hi


def _placeholder_png(w=8, h=8) -> io.BytesIO:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGBA", (w, h), (0, 0, 0, 0)).save(buf, "PNG")
    buf.seek(0)
    return buf


def add_model3d(slide, item, x, y, cx, cy, base_dir: Path):
    """3D 모델 아이템. item 키 — src(glb) · rot [x,y,z]도 · camera {pos,up,lookAt,fov} ·
    trans {meterPerModelUnit,preTrans,scale,postTrans} · viewport(EMU) · poster(PNG) · name · descr."""
    src = Path(item["src"])
    if not src.is_absolute():
        src = base_dir / src
    if not src.exists():
        raise ValueError(f"model3d src 없음: {src}")
    glb = _glb_part(slide, src)
    glb_rid = slide.part.relate_to(glb, RT_MODEL3D)
    poster = item.get("poster")
    if poster:
        pp = Path(poster)
        if not pp.is_absolute():
            pp = base_dir / pp
        _, img_rid = slide.part.get_or_add_image_part(str(pp))
    else:   # 정지 그림이 없으면 투명 자리 — PowerPoint 는 glb 를 직접 그린다(구형 뷰어만 이 그림을 본다)
        _, img_rid = slide.part.get_or_add_image_part(_placeholder_png())
    rot = item.get("rot", [0, 0, 0])
    cam = {**CAMERA, **(item.get("camera") or {})}
    tr = {**TRANS, **(item.get("trans") or {})}
    vp = int(item.get("viewport") or round(cx * VIEWPORT_RATIO))
    sid = _next_shape_id(slide)
    name = item.get("name") or f"3D 모델 {sid - 1}"
    descr = item.get("descr", "")
    r = [int(round(float(v) * 60000)) for v in rot]
    sc = "".join(f'<am3d:s{k} n="{int(round(float(v) * 1000000))}" d="1000000"/>'
                 for k, v in zip("xyz", tr["scale"]))
    xml = (
        f'<mc:AlternateContent xmlns:mc="{NS_MC}" xmlns:p="{NS_P}" xmlns:a="{NS_A}" xmlns:r="{NS_R}">'
        f'<mc:Choice xmlns:am3d="{NS_AM3D}" Requires="am3d"><p:graphicFrame><p:nvGraphicFramePr>'
        f'<p:cNvPr id="{sid}" name="{name}" descr="{descr}"/><p:cNvGraphicFramePr>'
        f'<a:graphicFrameLocks noChangeAspect="1"/></p:cNvGraphicFramePr><p:nvPr/></p:nvGraphicFramePr>'
        f'<p:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></p:xfrm><a:graphic>'
        f'<a:graphicData uri="{NS_AM3D}"><am3d:model3d r:embed="{glb_rid}"><am3d:spPr><a:xfrm>'
        f'<a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/>'
        f'</a:prstGeom></am3d:spPr><am3d:camera><am3d:pos x="{cam["pos"][0]}" y="{cam["pos"][1]}" '
        f'z="{cam["pos"][2]}"/><am3d:up dx="{cam["up"][0]}" dy="{cam["up"][1]}" dz="{cam["up"][2]}"/>'
        f'<am3d:lookAt x="{cam["lookAt"][0]}" y="{cam["lookAt"][1]}" z="{cam["lookAt"][2]}"/>'
        f'<am3d:perspective fov="{cam["fov"]}"/></am3d:camera><am3d:trans>'
        f'<am3d:meterPerModelUnit n="{tr["meterPerModelUnit"][0]}" d="{tr["meterPerModelUnit"][1]}"/>'
        f'<am3d:preTrans dx="{tr["preTrans"][0]}" dy="{tr["preTrans"][1]}" dz="{tr["preTrans"][2]}"/>'
        f'<am3d:scale>{sc}</am3d:scale><am3d:rot ax="{r[0]}" ay="{r[1]}" az="{r[2]}"/>'
        f'<am3d:postTrans dx="{tr["postTrans"][0]}" dy="{tr["postTrans"][1]}" dz="{tr["postTrans"][2]}"/>'
        f'</am3d:trans><am3d:raster rName="Office3DRenderer" rVer="16.0.8326"><am3d:blip r:embed="{img_rid}"/>'
        f'</am3d:raster><am3d:objViewport viewportSz="{vp}"/>{item.get("lights_xml") or LIGHTS}'
        f'</am3d:model3d></a:graphicData></a:graphic></p:graphicFrame></mc:Choice>'
        f'<mc:Fallback><p:pic><p:nvPicPr><p:cNvPr id="{sid}" name="{name}" descr="{descr}"/><p:cNvPicPr>'
        f'<a:picLocks noGrp="1" noRot="1" noChangeAspect="1" noMove="1" noResize="1" noEditPoints="1" '
        f'noAdjustHandles="1" noChangeArrowheads="1" noChangeShapeType="1" noCrop="1"/></p:cNvPicPr><p:nvPr/>'
        f'</p:nvPicPr><p:blipFill><a:blip r:embed="{img_rid}"/><a:stretch><a:fillRect/></a:stretch>'
        f'</p:blipFill><p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic></mc:Fallback></mc:AlternateContent>')
    el = etree.fromstring(xml)
    slide.shapes._spTree.append(el)
    return _Handle(el, sid, name)


class _Handle:
    """render_page 의 모핑 이름·애니 대상 추적용 — python-pptx 도형 흉내(name·shape_id·_element)."""

    def __init__(self, el, sid, name):
        self._element = el
        self.shape_id = sid
        self._name = name

    @property
    def name(self):
        return self._name

    @name.setter
    def name(self, v):
        self._name = v
        for nv in self._element.iter(f"{{{NS_P}}}cNvPr"):
            nv.set("name", v)


MIME = {".mp4": "video/mp4", ".m4v": "video/mp4", ".mov": "video/quicktime", ".wmv": "video/x-ms-wmv",
        ".mp3": "audio/mpeg", ".m4a": "audio/mp4", ".wav": "audio/wav"}


def media_duration(path) -> float:
    """영상·소리 길이(초) — ffprobe. 못 재면 0."""
    import subprocess
    try:
        out = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0",
                              str(path)], capture_output=True, text=True, timeout=30).stdout.strip()
        return round(float(out), 3)
    except Exception:
        return 0.0


def fill_media_dur(item, base_dir: Path):
    """media-play 에 dur 가 없으면 재생 길이(자르기 반영)로 채운다 — PowerPoint 도 cmd cTn dur 에 길이를 쓴다."""
    a = item.get("anim")
    if not a:
        return item
    lst = a if isinstance(a, list) else [a]
    if not any(isinstance(x, dict) and x.get("effect") == "media-play" and "dur" not in x for x in lst):
        return item
    src = Path(item["path"])
    if not src.is_absolute():
        src = base_dir / src
    d = media_duration(src)
    if item.get("trim"):
        st, en = item["trim"]
        d = (en or d) - (st or 0)
    new = [dict(x, dur=round(d, 3)) if isinstance(x, dict) and x.get("effect") == "media-play" and "dur" not in x
           else x for x in lst]
    return {**item, "anim": new if isinstance(a, list) else new[0]}


def add_media(slide, item, x, y, cx, cy, base_dir: Path):
    """영상·소리 아이템 — 개체만 만든다. 재생 timing 은 motion.media_entries 가 만든다.
    item 키 — path · poster(PNG, 첫 프레임 그림) · crop {l,t,r,b}(0~1) · trim [시작초, 끝초] ·
    fade [인초, 아웃초] · name."""
    from pptx.media import Video
    src = Path(item["path"])
    if not src.is_absolute():
        src = base_dir / src
    kind = "audio" if item["type"] == "audio" else "video"
    mime = MIME.get(src.suffix.lower(), "video/mp4" if kind == "video" else "audio/mpeg")
    media = Video.from_path_or_file_like(str(src), mime)
    part = slide.part._package.get_or_add_media_part(media)
    media_rid = slide.part.relate_to(part, RT.MEDIA)
    link_rid = slide.part.relate_to(part, RT.VIDEO if kind == "video" else RT.AUDIO)
    poster = item.get("poster")
    if poster:
        pp = Path(poster)
        if not pp.is_absolute():
            pp = base_dir / pp
        _, img_rid = slide.part.get_or_add_image_part(str(pp))
    else:
        _, img_rid = slide.part.get_or_add_image_part(_placeholder_png(16, 9))
    sid = _next_shape_id(slide)
    name = item.get("name") or src.name
    crop = ""
    if item.get("crop"):
        c = item["crop"]
        crop = "<a:srcRect " + " ".join(f'{k}="{int(round(float(c[k]) * 100000))}"' for k in "ltrb"
                                         if c.get(k)) + "/>"
    ext = f'<p14:media xmlns:p14="{NS_P14}" r:embed="{media_rid}"'
    inner = ""

    def ms(v):      # 실물은 소수 ms 를 쓴다(trim end="2564.3813")
        return f"{float(v) * 1000:.4f}".rstrip("0").rstrip(".")

    if item.get("trim"):
        st, en = item["trim"]
        inner += "<p14:trim" + (f' st="{ms(st)}"' if st else "") + (f' end="{ms(en)}"' if en else "") + "/>"
    if item.get("fade"):
        fi, fo = item["fade"]
        inner += "<p14:fade" + (f' in="{ms(fi)}"' if fi else "") + (f' out="{ms(fo)}"' if fo else "") + "/>"
    ext += f">{inner}</p14:media>" if inner else "/>"
    tag = "videoFile" if kind == "video" else "audioFile"
    xml = (f'<p:pic xmlns:p="{NS_P}" xmlns:a="{NS_A}" xmlns:r="{NS_R}"><p:nvPicPr>'
           f'<p:cNvPr id="{sid}" name="{name}"><a:hlinkClick r:id="" action="ppaction://media"/></p:cNvPr>'
           f'<p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr><a:{tag} r:link="{link_rid}"/>'
           f'<p:extLst><p:ext uri="{{DAA4B4D4-6D71-4841-9C94-3DE7FCFB9230}}">{ext}</p:ext></p:extLst>'
           f'</p:nvPr></p:nvPicPr><p:blipFill><a:blip r:embed="{img_rid}"/>{crop}<a:stretch/></p:blipFill>'
           f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
           f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>')
    el = etree.fromstring(xml)
    slide.shapes._spTree.append(el)
    return _Handle(el, sid, name)


def set_crop(pic, crop: dict):
    """그림 자르기 a:srcRect — crop {l,t,r,b} 0~1(음수 = 바깥 여백). 모핑이 자르기를 보간한다."""
    blip_fill = pic._element.find(f"{{{NS_P}}}blipFill")
    old = blip_fill.find(f"{{{NS_A}}}srcRect")
    if old is not None:
        blip_fill.remove(old)
    sr = etree.SubElement(blip_fill, f"{{{NS_A}}}srcRect")
    for k in "ltrb":
        if crop.get(k):
            sr.set(k, str(int(round(float(crop[k]) * 100000))))
    blip_fill.find(f"{{{NS_A}}}blip").addnext(sr)


def set_flip(shape, h=False, v=False):
    x = shape._element.find(f".//{{{NS_A}}}xfrm")
    if h:
        x.set("flipH", "1")
    if v:
        x.set("flipV", "1")


def add_raw(slide, item):
    """원본 개체 XML 그대로(`raw` 아이템) — 글상자처럼 관계(rId)가 없는 p:sp, 또는 그런 p:sp 만 담은 p:grpSp.

    줄마다 글자 크기가 다르게 바뀌는 글상자 모핑(제목 24·24·54pt → 28pt×3)은 그림으로
    옮길 수 없다. 원본 글상자를 테마 색·글꼴을 실제 값으로 풀어 그대로 붙이면 PowerPoint 가 원본과
    똑같이 줄마다 보간한다. id 만 이 장에서 새로 매긴다.
    """
    el = etree.fromstring(item["xml"])
    if el.xpath(".//@r:id|.//@r:embed|.//@r:link", namespaces={"r": NS_R}):
        raise ValueError("raw 아이템은 관계(rId)를 가진 개체를 받지 않는다 — 그림·영상은 image·video 아이템으로")
    sid = _next_shape_id(slide)
    nvs = list(el.iter(f"{{{NS_P}}}cNvPr"))
    for k, nv in enumerate(nvs):          # 묶음(grpSp)이면 안쪽 개체 번호도 이 장에서 새로 — 겹치면 PowerPoint 복구
        nv.set("id", str(sid + k))
    if item.get("name"):
        nvs[0].set("name", str(item["name"]))
    slide.shapes._spTree.append(el)
    return el
